"""Extracts text and images from a .pptx file, slide by slide.

Text is pulled directly from shapes (accurate, no OCR needed).
Images are saved to disk so the VLM can interpret them separately.
"""

import uuid
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.core.config import settings
from app.models.schemas import SlideImage, SlidePage


def _extract_slide_text(slide) -> str:
    lines = []
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            lines.append(shape.text_frame.text.strip())
        if shape.has_table:
            for row in shape.table.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            lines.append(f"[speaker notes] {notes}")
    return "\n".join(lines)


def _extract_slide_images(slide, deck_dir: Path, page_index: int) -> list[SlideImage]:
    images = []
    for shape in slide.shapes:
        if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        image = shape.image
        image_id = f"p{page_index:03d}_{uuid.uuid4().hex[:8]}"
        file_path = deck_dir / f"{image_id}.{image.ext}"
        file_path.write_bytes(image.blob)
        images.append(SlideImage(image_id=image_id, file_path=str(file_path)))
    return images


def extract_deck(pptx_path: Path, deck_id: str) -> list[SlidePage]:
    deck_dir = settings.upload_dir / deck_id / "images"
    deck_dir.mkdir(parents=True, exist_ok=True)

    presentation = Presentation(pptx_path)
    pages: list[SlidePage] = []
    for page_index, slide in enumerate(presentation.slides):
        raw_text = _extract_slide_text(slide)
        images = _extract_slide_images(slide, deck_dir, page_index)
        pages.append(SlidePage(page_index=page_index, raw_text=raw_text, images=images))
    return pages
