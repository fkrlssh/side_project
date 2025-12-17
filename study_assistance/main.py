"""Study assistance CLI for summarizing PPTX decks and generating quizzes."""
from __future__ import annotations

import argparse
import os
from pathlib import Path
from typing import Iterable, List, Sequence, Tuple

from fpdf import FPDF
from openai import OpenAI
from pptx import Presentation


def extract_slide_text(pptx_path: Path) -> List[str]:
    """Extract slide text while skipping empty slides."""
    presentation = Presentation(pptx_path)
    slides: List[str] = []
    for slide in presentation.slides:
        pieces: List[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()
                if text:
                    pieces.append(text)
        slide_text = "\n".join(pieces).strip()
        if slide_text:
            slides.append(slide_text)
    return slides


def build_summary_prompt(slides: Sequence[str]) -> str:
    numbered = [f"- 슬라이드 {idx}:\n{text}" for idx, text in enumerate(slides, start=1)]
    slide_blob = "\n\n".join(numbered)
    return (
        "당신은 학습 보조 전문가입니다. 아래 강의 슬라이드 내용을 한국어로 요약해 주세요. "
        "슬라이드별 핵심을 bullet point로 정리하고, 마지막에 '핵심 정리' 섹션을 별도로 추가하세요.\n\n"
        f"슬라이드 내용:\n{slide_blob}"
    )


def build_quiz_prompt(slides: Sequence[str], question_count: int) -> str:
    slide_blob = "\n\n".join(slides)
    return (
        "아래 강의 내용을 바탕으로 한국어 객관식 연습 문제를 만들어 주세요. "
        f"총 {question_count}문항을 생성하며, 각 문항마다 질문, 4개의 선택지, 정답 선택지 글자, "
        "짧은 해설을 포함해 주세요.\n\n"
        f"강의 노트:\n{slide_blob}"
    )


def run_llm(prompt: str, model: str, api_key: str | None = None) -> str:
    api_key = api_key or os.environ.get("OPENAI_API_KEY")
    if not api_key:
        raise ValueError("OPENAI_API_KEY가 필요합니다. 환경 변수 또는 --api-key 옵션을 사용하세요.")

    client = OpenAI(api_key=api_key)
    response = client.chat.completions.create(
        model=model,
        messages=[{"role": "user", "content": prompt}],
        temperature=0.5,
    )
    return response.choices[0].message.content.strip()


def write_pdf(title: str, sections: Iterable[Tuple[str, str]], output_path: Path) -> None:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)

    pdf.add_page()
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 10, title, ln=True)
    pdf.ln(6)

    for heading, body in sections:
        pdf.set_font("Helvetica", "B", 14)
        pdf.multi_cell(0, 8, heading)
        pdf.ln(2)
        pdf.set_font("Helvetica", size=12)
        pdf.multi_cell(0, 7, body)
        pdf.ln(4)

    pdf.output(str(output_path))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="PPTX 강의 자료를 요약하고 연습 문제 PDF를 생성하는 도구",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument("pptx", type=Path, help="입력 PPTX 파일 경로")
    parser.add_argument("--summary-pdf", type=Path, default=Path("summary.pdf"), help="요약 PDF 출력 경로")
    parser.add_argument("--quiz-pdf", type=Path, help="퀴즈 PDF 출력 경로(생략하면 생성 안 함)")
    parser.add_argument("--model", default="gpt-4o-mini", help="OpenAI LLM 모델명")
    parser.add_argument("--api-key", help="OpenAI API 키 (환경 변수 OPENAI_API_KEY 대신 지정 가능)")
    parser.add_argument("--questions", type=int, default=5, help="생성할 퀴즈 문항 수")
    return parser.parse_args()


def main() -> None:
    args = parse_args()

    if not args.pptx.exists():
        raise FileNotFoundError(f"PPTX 파일을 찾을 수 없습니다: {args.pptx}")

    slides = extract_slide_text(args.pptx)
    if not slides:
        raise ValueError("슬라이드에서 추출된 텍스트가 없습니다. PPTX 내용을 확인하세요.")

    summary_prompt = build_summary_prompt(slides)
    summary_text = run_llm(summary_prompt, model=args.model, api_key=args.api_key)
    write_pdf("요약 정리", [("요약", summary_text)], args.summary_pdf)

    if args.quiz_pdf:
        quiz_prompt = build_quiz_prompt(slides, args.questions)
        quiz_text = run_llm(quiz_prompt, model=args.model, api_key=args.api_key)
        write_pdf("연습 문제", [("문제 / 정답 / 해설", quiz_text)], args.quiz_pdf)

    print(f"요약 PDF 생성 완료: {args.summary_pdf}")
    if args.quiz_pdf:
        print(f"퀴즈 PDF 생성 완료: {args.quiz_pdf}")


if __name__ == "__main__":
    main()
